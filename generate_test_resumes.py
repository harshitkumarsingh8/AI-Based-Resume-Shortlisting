from __future__ import annotations

from pathlib import Path

try:
    from docx import Document
except ImportError:
    raise SystemExit(
        "python-docx is missing. Install it using:\n"
        "python -m pip install python-docx"
    )


# ============================================================
# CONFIGURATION
# ============================================================

PROJECT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = PROJECT_DIR / "benchmark_resumes"

TOTAL_SYNTHETIC_RESUMES = 96
GENERATED_PREFIX = "synthetic_resume_"


# These names and profiles are fictional test data.
FIRST_NAMES = [
    "Aarav", "Ananya", "Rohan", "Priya",
    "Ishaan", "Neha", "Kabir", "Meera",
    "Arjun", "Sana", "Dev", "Kavya",
]

LAST_NAMES = [
    "Sharma", "Verma", "Singh", "Gupta",
    "Mehta", "Kumar", "Patel", "Reddy",
]

INSTITUTES = [
    "Example Institute of Technology",
    "Sample College of Computer Applications",
    "Demo Institute of Science",
    "Fictional University of Technology",
    "Example School of Computing",
    "Sample Technical University",
    "Demo College of Engineering",
    "Illustrative Institute of Information Technology",
]

DEGREES = [
    "Bachelor of Computer Applications (BCA)",
    "B.Tech in Computer Science",
    "B.Sc. in Computer Science",
    "B.Tech in Information Technology",
]

VARIANT_EXTRAS = [
    ["Excel", "Data Analysis"],
    ["Flask", "REST API"],
    ["Power BI", "Data Visualization"],
    ["Linux", "GitHub"],
    ["Statistics", "NumPy"],
    ["SQL", "Database Design"],
    ["Docker", "API Development"],
    ["Problem Solving", "Software Testing"],
]


# ============================================================
# SYNTHETIC PROFILE TEMPLATES
# ============================================================

PROFILES = [
    {
        "role": "AI/ML Engineer Intern",
        "summary": (
            "Fictional candidate interested in building and evaluating "
            "machine learning applications using structured datasets."
        ),
        "skills": [
            "Python", "Machine Learning", "Scikit-learn",
            "Pandas", "NumPy", "SQL", "Git",
            "Feature Engineering", "Model Evaluation",
        ],
        "project": (
            "Student Performance Prediction",
            "Prepared a sample dataset with Pandas and trained a "
            "baseline classification model using Scikit-learn.",
        ),
    },
    {
        "role": "Data Science Intern",
        "summary": (
            "Fictional candidate focused on data cleaning, statistical "
            "analysis, visualization, and predictive modelling."
        ),
        "skills": [
            "Python", "Data Science", "Pandas", "NumPy",
            "Statistics", "Data Visualization", "SQL",
            "Scikit-learn", "Exploratory Data Analysis",
        ],
        "project": (
            "Exploratory Data Analysis Dashboard",
            "Cleaned a sample dataset, analysed trends, and summarized "
            "observations using Python and visualizations.",
        ),
    },
    {
        "role": "Python Backend Developer",
        "summary": (
            "Fictional candidate interested in Python application "
            "development, API integration, and database operations."
        ),
        "skills": [
            "Python", "SQL", "FastAPI", "Flask",
            "REST API", "Git", "GitHub",
            "Database Design", "Streamlit",
        ],
        "project": (
            "Task Management API",
            "Created a sample API workflow with CRUD operations and "
            "database interaction.",
        ),
    },
    {
        "role": "Java Developer",
        "summary": (
            "Fictional candidate focused on object-oriented programming "
            "and backend application development."
        ),
        "skills": [
            "Java", "SQL", "MySQL", "Git",
            "GitHub", "Data Structures",
            "Algorithms", "Object-Oriented Programming",
            "REST API",
        ],
        "project": (
            "Library Management System",
            "Designed a sample application to maintain books, members, "
            "and borrowing records.",
        ),
    },
    {
        "role": "NLP Engineer Intern",
        "summary": (
            "Fictional candidate exploring text processing, information "
            "extraction, and semantic document matching."
        ),
        "skills": [
            "Python", "Natural Language Processing", "NLP",
            "Text Processing", "TF-IDF", "Text Similarity",
            "Semantic Embeddings", "Scikit-learn",
            "Sentence Transformers",
        ],
        "project": (
            "Document Similarity Prototype",
            "Compared sample text documents using TF-IDF and a "
            "similarity score for experimentation.",
        ),
    },
    {
        "role": "Data Analyst Intern",
        "summary": (
            "Fictional candidate interested in reporting, data quality, "
            "SQL queries, and business insights."
        ),
        "skills": [
            "SQL", "Python", "Microsoft Excel",
            "Data Analysis", "Data Cleaning",
            "Data Visualization", "Power BI",
            "Statistics", "Pandas",
        ],
        "project": (
            "Sales Reporting Dashboard",
            "Summarized fictional sales records and presented trends "
            "through tables and charts.",
        ),
    },
    {
        "role": "Software Engineering Fresher",
        "summary": (
            "Fictional recent graduate with academic programming "
            "experience and an interest in software development."
        ),
        "skills": [
            "Python", "Java", "Data Structures",
            "Algorithms", "SQL", "Git",
            "Problem Solving", "Software Testing",
        ],
        "project": (
            "Student Management System",
            "Implemented basic record creation, searching, updating, "
            "and deletion for sample student data.",
        ),
    },
    {
        "role": "Cloud and DevOps Intern",
        "summary": (
            "Fictional candidate learning deployment workflows, "
            "developer tools, and application infrastructure."
        ),
        "skills": [
            "Python", "Linux", "Git", "GitHub",
            "Docker", "Bash", "CI/CD",
            "Cloud Computing", "REST API",
        ],
        "project": (
            "Application Deployment Exercise",
            "Documented a sample application's setup, test execution, "
            "and container-based deployment workflow.",
        ),
    },
    {
        "role": "Frontend Developer Intern",
        "summary": (
            "Fictional candidate interested in building user interfaces "
            "and interactive web applications."
        ),
        "skills": [
            "JavaScript", "HTML", "CSS",
            "React", "Git", "GitHub",
            "REST API", "Problem Solving",
        ],
        "project": (
            "Interactive Portfolio Website",
            "Built a sample responsive interface with project sections "
            "and reusable UI components.",
        ),
    },
    {
        "role": "Database Developer Intern",
        "summary": (
            "Fictional candidate interested in relational data models, "
            "SQL querying, and database-backed software."
        ),
        "skills": [
            "SQL", "MySQL", "PostgreSQL",
            "Database Design", "Data Modeling",
            "Python", "Git", "REST API",
        ],
        "project": (
            "Inventory Database",
            "Designed sample tables and queries for tracking products "
            "and inventory transactions.",
        ),
    },
    {
        "role": "Generative AI Intern",
        "summary": (
            "Fictional candidate exploring language models, prompt "
            "design, and AI-assisted application development."
        ),
        "skills": [
            "Python", "Generative AI", "Large Language Models",
            "Prompt Engineering", "NLP",
            "Semantic Embeddings", "Streamlit",
            "Git", "Information Extraction",
        ],
        "project": (
            "Knowledge Assistant Prototype",
            "Created a conceptual interface for submitting questions "
            "and presenting responses from sample text.",
        ),
    },
    {
        "role": "Machine Learning Research Intern",
        "summary": (
            "Fictional candidate interested in experimentation, model "
            "evaluation, and comparing machine learning approaches."
        ),
        "skills": [
            "Python", "Machine Learning", "Deep Learning",
            "Scikit-learn", "TensorFlow",
            "Model Evaluation", "Statistics",
            "Feature Engineering", "NumPy",
        ],
        "project": (
            "Model Comparison Experiment",
            "Compared baseline model outputs on a sample dataset and "
            "documented evaluation observations.",
        ),
    },
]


# ============================================================
# DOCUMENT GENERATOR
# ============================================================

def add_bullet(document, text: str) -> None:
    """Add a readable bullet point to a DOCX document."""
    document.add_paragraph(text, style="List Bullet")


def create_resume(
    number: int,
    profile: dict,
) -> Path:
    """Create one fictional resume suitable for parser testing."""
    first_name = FIRST_NAMES[(number - 1) % len(FIRST_NAMES)]
    last_name = LAST_NAMES[
        ((number - 1) // len(FIRST_NAMES)) % len(LAST_NAMES)
    ]

    full_name = f"{first_name} {last_name}"
    candidate_id = f"SYNTH-{number:03d}"

    variant = (number - 1) // len(PROFILES)
    institute = INSTITUTES[number % len(INSTITUTES)]
    degree = DEGREES[number % len(DEGREES)]
    graduation_year = 2022 + (number % 5)

    skills = list(profile["skills"])
    extras = VARIANT_EXTRAS[(number - 1) % len(VARIANT_EXTRAS)]

    for skill in extras:
        if skill.casefold() not in {
            existing.casefold() for existing in skills
        }:
            skills.append(skill)

    document = Document()

    document.add_heading(full_name, level=0)
    document.add_paragraph(
        f"SYNTHETIC TEST RESUME | Candidate ID: {candidate_id}"
    )
    document.add_paragraph(
        f"Target Role: {profile['role']}"
    )
    document.add_paragraph(
        f"Email: {candidate_id.lower()}@example.com"
    )
    document.add_paragraph(
        "Phone: 000-000-0000 | Location: Example City, India"
    )

    document.add_heading("Professional Summary", level=1)
    document.add_paragraph(profile["summary"])
    document.add_paragraph(
        f"This fictional profile is template variation {variant + 1}. "
        "All details are synthetic and created for software testing."
    )

    document.add_heading("Technical Skills", level=1)
    document.add_paragraph(", ".join(skills))

    document.add_heading("Education", level=1)
    document.add_paragraph(
        f"{degree} — {institute} | {graduation_year}"
    )

    document.add_heading("Experience", level=1)

    if number % 3 == 0:
        document.add_paragraph(
            "Experience Level: Fictional internship experience."
        )
        add_bullet(
            document,
            "Assisted with a sample technical task and documented "
            "the implementation steps.",
        )
    else:
        document.add_paragraph(
            "Experience Level: Fresher / academic project experience."
        )
        add_bullet(
            document,
            "Practised programming and data handling through "
            "fictional coursework exercises.",
        )

    document.add_heading("Projects", level=1)

    project_name, project_description = profile["project"]
    document.add_paragraph(project_name)
    add_bullet(document, project_description)
    add_bullet(
        document,
        f"Used selected technologies including {', '.join(skills[:4])}.",
    )

    document.add_heading("Certifications", level=1)
    document.add_paragraph(
        "Synthetic example entries: Programming Fundamentals; "
        "Introduction to Data and Software Tools."
    )

    document.add_heading("Testing Notice", level=1)
    document.add_paragraph(
        "This is artificial benchmark data, not a real person's resume. "
        "It must not be used for actual recruitment decisions."
    )

    file_path = OUTPUT_DIR / (
        f"{GENERATED_PREFIX}{number:03d}.docx"
    )

    document.save(file_path)
    return file_path


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("=" * 60)
    print("SYNTHETIC RESUME DATASET GENERATOR")
    print("=" * 60)
    print(f"Output folder: {OUTPUT_DIR}")
    print(f"Synthetic resumes to create: {TOTAL_SYNTHETIC_RESUMES}")

    created = 0

    for number in range(1, TOTAL_SYNTHETIC_RESUMES + 1):
        profile = PROFILES[(number - 1) % len(PROFILES)]

        file_path = create_resume(
            number=number,
            profile=profile,
        )

        created += 1

        if number % 12 == 0:
            print(
                f"Created {number}/{TOTAL_SYNTHETIC_RESUMES} "
                "synthetic resumes..."
            )

    print("\n" + "=" * 60)
    print("GENERATION COMPLETE")
    print("=" * 60)
    print(f"New synthetic files created: {created}")
    print(f"Output folder: {OUTPUT_DIR}")
    print(
        "Existing resumes were not deleted. "
        "Only synthetic_resume_*.docx files were generated/overwritten."
    )
    print(
        "\nNOTE: These resumes use varied fictional role profiles, "
        "but they are template-generated. They are useful for pipeline "
        "and performance tests, not for measuring real-world accuracy."
    )


if __name__ == "__main__":
    main()