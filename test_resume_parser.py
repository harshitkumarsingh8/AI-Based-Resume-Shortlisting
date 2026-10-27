
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from utils.resume_parser import (
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_resume_text,
)


def test_unsupported_file_type():
    try:
        extract_resume_text("resume.txt")
    except ValueError as error:
        assert "Unsupported file type" in str(error)
    else:
        raise AssertionError("Expected ValueError for TXT file")


def test_parser_functions_exist():
    assert callable(extract_text_from_pdf)
    assert callable(extract_text_from_docx)
    assert callable(extract_resume_text)


def test_sample_pdf_text_extraction():
    pdf_path = Path("sample_resumes/sample_resume.pdf")

    assert pdf_path.exists(), "Sample PDF not found"

    text = extract_resume_text(pdf_path)

    assert "Aarav Sharma" in text
    assert "Python" in text
    assert "SQL" in text
    assert "PROJECTS" in text


def test_docx_text_extraction():
    with TemporaryDirectory() as temp_dir:
        docx_path = Path(temp_dir) / "test_resume.docx"

        document = Document()
        document.add_paragraph("Candidate: Rahul Kumar")
        document.add_paragraph("Skills: Python, SQL")
        document.save(docx_path)

        text = extract_text_from_docx(docx_path)

        assert "Rahul Kumar" in text
        assert "Python" in text
        assert "SQL" in text


def test_docx_resume_through_main_function():
    with TemporaryDirectory() as temp_dir:
        docx_path = Path(temp_dir) / "test_resume.docx"

        document = Document()
        document.add_paragraph("Education: BCA")
        document.save(docx_path)

        text = extract_resume_text(docx_path)

        assert "Education: BCA" in text
