from pathlib import Path
from docx import Document


# Project paths
PROJECT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = PROJECT_DIR / "benchmark_resumes"
OUTPUT_FILE = OUTPUT_DIR / "positive_test_resume.docx"


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # Resume heading
    doc.add_heading("Positive Test Candidate", level=0)
    doc.add_paragraph("Email: positive.test@example.com")
    doc.add_paragraph("Phone: 0000000000")
    doc.add_paragraph("Location: Test City")

    # Professional summary
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(
        "A fictional AI/ML candidate profile created exclusively "
        "for testing an AI-powered resume shortlisting system. "
        "Knowledge of Python, machine learning, deep learning, "
        "artificial intelligence, data analysis, and AI application "
        "development."
    )

    # Skills
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(
        "Python, Machine Learning, Deep Learning, "
        "Artificial Intelligence, SQL, Pandas, NumPy, "
        "Git, Streamlit"
    )

    # Education
    doc.add_heading("Education", level=1)
    doc.add_paragraph(
        "Bachelor's degree in Computer Science — "
        "Fictional test profile"
    )

    # Experience
    doc.add_heading("Project Experience", level=1)

    doc.add_heading("Machine Learning Project", level=2)
    doc.add_paragraph(
        "Developed a test machine learning workflow using Python, "
        "Pandas, and NumPy for data processing. Applied machine "
        "learning and deep learning concepts to an artificial "
        "intelligence use case."
    )

    doc.add_heading("Database and Application Project", level=2)
    doc.add_paragraph(
        "Used SQL for querying structured data, Git for version "
        "control, and Streamlit to create a prototype interface "
        "for presenting model results."
    )

    # Certifications
    doc.add_heading("Certifications", level=1)
    doc.add_paragraph(
        "No actual certifications claimed. This document is "
        "synthetic test data and is not a genuine job application."
    )

    # Save document
    doc.save(OUTPUT_FILE)

    print("Positive test resume created successfully.")
    print(f"File: {OUTPUT_FILE}")
    print("Type: Synthetic test data")
    print("Configured skills included:")
    print(
        "Python, Machine Learning, Deep Learning, SQL, Pandas, "
        "NumPy, Artificial Intelligence, Git, Streamlit"
    )


if __name__ == "__main__":
    main()