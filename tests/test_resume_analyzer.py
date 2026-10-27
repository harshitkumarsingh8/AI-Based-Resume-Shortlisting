
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from services.resume_analyzer import analyze_resume


SAMPLE_PDF = Path("sample_resumes/sample_resume.pdf")


def create_test_resume(file_path: Path, paragraphs: list[str]) -> None:
    """Create a controlled DOCX resume for testing."""
    document = Document()

    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    document.save(file_path)


def test_analyze_docx_resume():
    with TemporaryDirectory() as temp_dir:
        resume_path = Path(temp_dir) / "sample_resume.docx"

        create_test_resume(
            resume_path,
            [
                "Rahul Kumar",
                "rahul.kumar@example.com",
                "Phone: +91 9876543210",
                "EDUCATION",
                "Bachelor of Computer Applications (BCA), 2025",
                "EXPERIENCE",
                "Python Intern at ABC Technologies",
                "PROJECTS",
                "Student Performance Prediction using Python",
                "CERTIFICATIONS",
                "Python Programming Certificate",
                "SKILLS",
                "Python, SQL, Pandas and Machine Learning",
            ],
        )

        result = analyze_resume(resume_path)

        assert result["filename"] == "sample_resume.docx"
        assert "Rahul Kumar" in result["text"]
        assert "python" in result["skills"]
        assert "sql" in result["skills"]
        assert "pandas" in result["skills"]
        assert "machine learning" in result["skills"]
        assert result["skill_count"] >= 4


def test_extract_contact_details():
    with TemporaryDirectory() as temp_dir:
        resume_path = Path(temp_dir) / "contact_resume.docx"

        create_test_resume(
            resume_path,
            [
                "Rahul Kumar",
                "rahul.kumar@example.com",
                "Phone: +91 9876543210",
                "Skills: Python and SQL",
            ],
        )

        result = analyze_resume(resume_path)

        assert result["candidate_name"] == "Rahul Kumar"
        assert result["email"] == "rahul.kumar@example.com"
        assert result["phone"] is not None
        assert "9876543210" in result["phone"]


def test_extract_resume_sections():
    with TemporaryDirectory() as temp_dir:
        resume_path = Path(temp_dir) / "sections_resume.docx"

        create_test_resume(
            resume_path,
            [
                "Candidate Name",
                "EDUCATION",
                "Bachelor of Computer Applications",
                "EXPERIENCE",
                "Software Developer Intern",
                "PROJECTS",
                "Resume Screening System",
                "CERTIFICATIONS",
                "Python Certificate",
            ],
        )

        result = analyze_resume(resume_path)

        assert any(
            "Bachelor of Computer Applications" in item
            for item in result["education"]
        )
        assert any(
            "Software Developer Intern" in item
            for item in result["experience"]
        )
        assert any(
            "Resume Screening System" in item
            for item in result["projects"]
        )
        assert any(
            "Python Certificate" in item
            for item in result["certifications"]
        )


def test_missing_contact_details_are_not_invented():
    with TemporaryDirectory() as temp_dir:
        resume_path = Path(temp_dir) / "minimal_resume.docx"

        create_test_resume(
            resume_path,
            [
                "Rahul Kumar",
                "Skills: Python and SQL",
            ],
        )

        result = analyze_resume(resume_path)

        assert result["email"] is None
        assert result["phone"] is None


def test_analyze_sample_pdf_resume():
    assert SAMPLE_PDF.is_file(), (
        f"Sample PDF not found at: {SAMPLE_PDF.resolve()}"
    )

    result = analyze_resume(SAMPLE_PDF)

    assert result["filename"] == SAMPLE_PDF.name
    assert result["text"].strip()
    assert "Aarav Sharma" in result["text"]
    assert "python" in result["skills"]
    assert "sql" in result["skills"]
    assert result["skill_count"] >= 2


def test_missing_resume_raises_error():
    with TemporaryDirectory() as temp_dir:
        missing_path = Path(temp_dir) / "missing.docx"

        try:
            analyze_resume(missing_path)
        except FileNotFoundError:
            pass
        else:
            raise AssertionError("Expected FileNotFoundError")


def test_empty_docx_raises_error():
    with TemporaryDirectory() as temp_dir:
        resume_path = Path(temp_dir) / "empty.docx"

        Document().save(resume_path)

        try:
            analyze_resume(resume_path)
        except ValueError as error:
            assert "No readable text" in str(error)
        else:
            raise AssertionError("Expected ValueError")
